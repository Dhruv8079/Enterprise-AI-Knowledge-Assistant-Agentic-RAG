"""DOCX loader built on ``python-docx``.

Paragraph text and table content are extracted separately. Because DOCX has
no intrinsic pagination, paragraphs are grouped into logical blocks and tables
are emitted as their own units, each with descriptive metadata.
"""

from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document

from loaders.base import DocumentLoadError, LoadedDocument, clean_text
from utils.logger import get_logger

logger = get_logger(__name__)

_PARAGRAPHS_PER_BLOCK: int = 25


def _table_to_text(rows: list[list[str]]) -> str:
    """Render a table as pipe-delimited text, preserving the header row."""
    lines: list[str] = []
    for row_index, row in enumerate(rows):
        cleaned = [cell.replace("\n", " ").strip() for cell in row]
        if not any(cleaned):
            continue
        lines.append(" | ".join(cleaned))
        if row_index == 0:
            lines.append("-" * 60)
    return "\n".join(lines)


def load_docx(path: Path, document_id: str, filename: str) -> LoadedDocument:
    """Extract paragraphs and tables from a DOCX file.

    Args:
        path: Path to the DOCX on disk.
        document_id: Stable document identifier.
        filename: Original file name.

    Returns:
        A :class:`LoadedDocument` containing paragraph blocks and tables.

    Raises:
        DocumentLoadError: If the file is not a valid DOCX.
    """
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentLoadError(
            "python-docx is not installed. Run 'pip install python-docx' to enable DOCX support."
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentLoadError(f"Could not open DOCX '{filename}': {exc}") from exc

    documents: list[Document] = []

    paragraphs = [clean_text(p.text) for p in document.paragraphs]
    paragraphs = [p for p in paragraphs if p]
    for block_index in range(0, len(paragraphs), _PARAGRAPHS_PER_BLOCK):
        block = paragraphs[block_index : block_index + _PARAGRAPHS_PER_BLOCK]
        text = "\n".join(block)
        documents.append(
            Document(
                page_content=text,
                metadata={
                    "source": filename,
                    "file_type": "DOCX",
                    "page": f"paragraphs {block_index + 1}-{block_index + len(block)}",
                    "document_id": document_id,
                },
            )
        )

    for table_index, table in enumerate(document.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        text = clean_text(_table_to_text(rows))
        if not text:
            continue
        documents.append(
            Document(
                page_content=f"Table {table_index + 1} in {filename}:\n{text}",
                metadata={
                    "source": filename,
                    "file_type": "DOCX",
                    "page": f"table {table_index + 1}",
                    "element": "table",
                    "document_id": document_id,
                },
            )
        )

    if not documents:
        raise DocumentLoadError(f"No readable content found in '{filename}'.")

    logger.info("Extracted %d block(s)/table(s) from '%s'", len(documents), filename)
    return LoadedDocument(
        document_id=document_id,
        filename=filename,
        file_type="DOCX",
        size_bytes=path.stat().st_size,
        num_units=len(documents),
        documents=documents,
    )
